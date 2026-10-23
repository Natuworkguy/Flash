"""Designs: pages, app screens, one-pagers, social posts and decks that
Flash builds as HTML and the web UI shows in its side panel to work on.

In the panel the user comments on any part of a design, edits its words
in place, and drags the sliders Flash gave it. Words and slider values
are written back into the HTML file here; comments go to Flash as a
message. A design exports as a PNG, a PDF or, a deck, as slides.

A design system is a Markdown file in ~/.flash/design-systems that Flash
writes from a project's code or a brand file: its colours, type,
spacing and components, so every design made with it matches.
"""

import html
import re
from pathlib import Path
from typing import Optional

from . import workspace
from .paths import FLASH_DIR

SYSTEMS_DIR = FLASH_DIR / "design-systems"
MAX_EDITS = 200
MAX_EDIT_CHARS = 5000
EXPORT_FORMATS = ("png", "pdf", "pptx")
SLIDE_SIZE = (1920, 1080)
EXPORT_WIDTH = 1440
NAVIGATION_TIMEOUT_MS = 20_000
SETTLE_MS = 400

# A custom property as written in a stylesheet: its name, then its value
# up to the end of the declaration.
_VAR = r"(--{name}\s*:\s*)([^;{{}}]+?)(\s*[;}}])"
_SAFE_VALUE = re.compile(r"^[^;{}<>]{1,120}$")
_VAR_NAME = re.compile(r"^--[A-Za-z0-9_-]{1,60}$")

# What the model is told about designs.
PROMPT = (Path(__file__).parent / "design_prompt.txt").read_text(
    encoding="utf-8"
).strip()


def systems() -> list[dict]:
    """The design systems, each {"name", "path", "about"}: its file
    name, where it is, and its first line of text."""

    found = []
    for path in sorted(SYSTEMS_DIR.glob("*.md")):
        try:
            lines = path.read_text(encoding="utf-8").splitlines()
        except (OSError, UnicodeDecodeError):
            continue
        about = next(
            (ln.strip("# ").strip() for ln in lines if ln.strip()), "",
        )
        found.append({"name": path.stem, "path": str(path), "about": about})
    return found


def prompt_block() -> str:
    """How to make designs, and the design systems there are."""

    listed = systems()
    if not listed:
        return PROMPT
    names = "\n".join(
        f"- {s['name']}: {s['about'][:120]} ({s['path']})" for s in listed
    )
    return f"{PROMPT}\nThe design systems saved so far:\n{names}"


class DesignError(ValueError):
    pass


def _replace_text(source: str, old: str, new: str) -> Optional[str]:
    """SOURCE with the words OLD changed to NEW where they first appear
    as the text of an element, or None when they cannot be found."""

    for written in (html.escape(old, quote=False), old):
        # Between tags, so a word in a class name or a script is safe.
        at = re.search(
            r"(>\s*)" + re.escape(written) + r"(\s*<)", source,
        )
        if at:
            start, end = at.start() + len(at.group(1)), at.end() - len(
                at.group(2))
            return source[:start] + html.escape(new, quote=False) + \
                source[end:]
    return None


def _set_var(source: str, name: str, value: str) -> Optional[str]:
    pattern = re.compile(_VAR.format(name=re.escape(name[2:])))
    if not pattern.search(source):
        return None
    return pattern.sub(lambda m: m.group(1) + value + m.group(3), source, 1)


def apply(source: str, edits: list, values: dict) -> tuple[str, list]:
    """SOURCE with the user's word EDITS ({"old", "new"} each) and
    slider VALUES (custom property name to value) written in. Returns
    the new text and the edits that could not be placed."""

    missed = []
    for edit in list(edits)[:MAX_EDITS]:
        old = str((edit or {}).get("old") or "").strip()
        new = str((edit or {}).get("new") or "").strip()
        if not old or old == new or len(new) > MAX_EDIT_CHARS:
            continue
        changed = _replace_text(source, old, new)
        if changed is None:
            missed.append({"old": old, "new": new})
        else:
            source = changed
    for name, value in dict(values or {}).items():
        name, value = str(name), str(value).strip()
        if not _VAR_NAME.match(name) or not _SAFE_VALUE.match(value):
            continue
        source = _set_var(source, name, value) or source
    return source, missed


def save(file_id: str, edits: list, values: dict) -> dict:
    """Write the user's edits into a shown page, its kept copy and the
    file it came from. Returns {"path", "missed"}."""

    kept = workspace.kept_file(file_id)
    if kept is None or kept[1] != "text/html":
        raise DesignError("no such design")
    try:
        before = kept[0].read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        raise DesignError(f"could not read the design: {exc}") from exc
    text, missed = apply(before, edits, values)
    if text == before:
        return {"path": "", "missed": missed}
    kept[0].write_text(text, encoding="utf-8")
    source = workspace.file_source(file_id)
    written = ""
    if source and Path(source).is_file():
        try:
            Path(source).write_text(text, encoding="utf-8")
            written = source
        except OSError as exc:
            raise DesignError(f"could not save {source}: {exc}") from exc
    return {"path": written, "missed": missed}


def export(file_id: str, fmt: str, out_dir: Path) -> Path:
    """The design exported as FMT (png, pdf or pptx) into OUT_DIR, drawn
    by the same headless Chromium Flash takes screenshots with."""

    if fmt not in EXPORT_FORMATS:
        raise DesignError(f"cannot export as {fmt}")
    kept = workspace.kept_file(file_id)
    if kept is None or kept[1] != "text/html":
        raise DesignError("no such design")
    try:
        from playwright.sync_api import Error as PlaywrightError
        from playwright.sync_api import sync_playwright

        from .browser import launch_chromium as browser_launch
    except ImportError as exc:
        raise DesignError(
            "Exporting needs Playwright: pip install playwright, then "
            "playwright install chromium"
        ) from exc
    if fmt == "pptx":
        try:
            import pptx  # noqa: F401  (only checking it is there)
        except ImportError as exc:
            raise DesignError(
                "Slides need python-pptx: pip install python-pptx"
            ) from exc

    name = Path(workspace.file_source(file_id) or kept[0]).stem
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{name}.{fmt}"
    try:
        with sync_playwright() as driver:
            browser = browser_launch(driver)
            try:
                page = browser.new_page(viewport={
                    "width": SLIDE_SIZE[0] if fmt == "pptx" else EXPORT_WIDTH,
                    "height": SLIDE_SIZE[1] if fmt == "pptx" else 900,
                })
                page.goto(kept[0].resolve().as_uri(), wait_until="load",
                          timeout=NAVIGATION_TIMEOUT_MS)
                page.wait_for_timeout(SETTLE_MS)
                if fmt == "png":
                    page.screenshot(path=str(out), full_page=True)
                elif fmt == "pdf":
                    page.emulate_media(media="screen")
                    page.pdf(path=str(out), print_background=True,
                             prefer_css_page_size=True)
                else:
                    shots = []
                    for i, slide in enumerate(
                            page.locator("section.slide").all()):
                        shot = out_dir / f"{name}-slide-{i + 1}.png"
                        slide.screenshot(path=str(shot))
                        shots.append(shot)
                    if not shots:
                        raise DesignError(
                            "This design has no slides: a deck is one "
                            "<section class=\"slide\"> per slide."
                        )
                    _slides(shots, out)
                    for shot in shots:
                        shot.unlink(missing_ok=True)
            finally:
                browser.close()
    except PlaywrightError as exc:
        raise DesignError(str(exc).strip().splitlines()[0]) from exc
    return out


def _slides(pictures: list, out: Path) -> None:
    """A .pptx with one 16:9 slide per picture, each filling its slide."""

    from pptx import Presentation
    from pptx.util import Emu

    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
    blank = deck.slide_layouts[6]
    for picture in pictures:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(picture), 0, 0, deck.slide_width, deck.slide_height,
        )
    deck.save(str(out))
